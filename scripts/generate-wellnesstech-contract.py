"""Generate WellnessTech website/platform development agreement as .docx"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.15


def set_run_font(run, bold=False, size=11, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, bold=True, size=14)
    p.paragraph_format.space_after = Pt(4)


def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, bold=True, size=11)
    p.paragraph_format.space_after = Pt(12)


def add_heading_custom(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=True, size=12)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def add_body(text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic, size=11)
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11)
    p.paragraph_format.left_indent = Inches(0.25)
    return p


def add_blank_line():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)


add_title("WEBSITE & DIGITAL PLATFORM DEVELOPMENT")
add_title("AND MANAGEMENT SERVICES AGREEMENT")
add_subtitle("(Independent Contractor Agreement — IT Professional)")

add_body(
    'This Website & Digital Platform Development and Management Services Agreement ("Agreement") '
    'is entered into as of _____________________ ("Effective Date"), by and between:'
)

add_body(
    "Client: WellnessTech Corporation, parent company of KIAN Privé, WellnessTech Bio Distribution, "
    "WellnessTech Longevity and Research Clinics, and related affiliated clinic brands (including Facial "
    'Aesthetics / Facial Design Studio, 4EverGlow Wellness, and JOA as and when developed) ("Client"), '
    "with a principal place of business at ____________________________________."
)

add_body(
    'Contractor: ____________________________________ ("Contractor"), an independent IT professional '
    "and/or firm with a principal place of business at ____________________________________."
)

add_body('Client and Contractor may each be referred to as a "Party" and collectively as the "Parties."')

add_heading_custom("1. Engagement & Scope of Services")
add_body(
    "1.1 Engagement. Client engages Contractor, as an independent contractor, to design, build, deploy, "
    "integrate, and/or maintain Client's websites and related digital platforms (collectively, the \"Platforms\"), "
    "including public marketing sites, clinical intake applications, partner-clinic sites, commerce/shop "
    "functionality, practitioner and admin portals, APIs, and related backend systems, as specified in Exhibit A "
    'and any subsequent Statements of Work ("SOWs").'
)
add_body(
    "1.2 Covered Properties. Unless otherwise limited in writing, Platforms may include properties for Client "
    "and/or its subsidiaries and affiliated brands, including without limitation:"
)
for item in [
    "KIAN Privé (primary brand website, shop, bookings, member/provider portals)",
    "KIAN Privé Wellness Hub / Privé Therapeutics (clinical intake and provider-connect experience)",
    "WellnessTech Bio Distribution (distribution / wholesale digital properties)",
    "Facial Aesthetics / Facial Design Studio",
    "4EverGlow Wellness",
    "JOA (as and when developed)",
    "Such other Client or affiliate digital properties listed in Exhibit A or an SOW",
]:
    add_bullet(item)

add_body('1.3 Services. Services under this Agreement ("Services") include, without limitation:')
for item in [
    "Design, development, and deployment of front-end and back-end functionality for the Platforms",
    "Custom application development (including Next.js, TanStack Start / similar modern web frameworks, APIs, and database layers)",
    "Configuration and integration of third-party tools (booking/scheduling, payment processing including Authorize.net or successor processors, email/transactional messaging, CRM, analytics)",
    "Clinical and wellness intake forms, signature capture, review-deposit / fee collection flows, and forwarding of submissions into Client's central charting / EHR workflow (including Wellness Tech EHR routing as directed by Client)",
    "Partner-site intake submission APIs and cross-site data handoff to Client systems",
    "Practitioner, ambassador, partner, and admin portal features as scoped",
    "Ongoing maintenance, including dependency/CMS/plugin updates where applicable, security patches, backups, and environment configuration",
    "Performance monitoring, uptime management, and troubleshooting of technical issues",
    "Implementation of content, feature, and design changes requested by Client",
    "Domain, hosting, SSL, DNS, and repository access management (as applicable and authorized)",
    "Basic SEO and technical compliance upkeep (accessibility, page speed)",
    "Such other services as the Parties agree in a written SOW or signed change order",
]:
    add_bullet(item)

add_body(
    "1.4 Statement of Work. A detailed description of the initial project scope, deliverables, specifications, "
    "repositories, environments, and timeline is set out in Exhibit A (Statement of Work), attached hereto and "
    "incorporated by reference. Any material change to scope, timeline, or fees must be documented in a written "
    "change order signed by both Parties before the changed work begins."
)
add_body(
    "1.5 Clinical / Business Decisions Out of Scope. Contractor provides technology Services only. Contractor does "
    "not practice medicine, prescribe medications, make clinical determinations, approve therapies, or provide legal, "
    "compliance, or medical advice. Client remains solely responsible for clinical content accuracy, prescribing "
    "authority, medical director oversight, pharmacy fulfillment decisions, and regulatory compliance of Client's "
    "operations."
)

add_heading_custom("2. Term")
add_body(
    "This Agreement begins on the Effective Date and continues for an initial term of ____________ months "
    '("Initial Term"), and will automatically renew for successive ____________-month periods unless either Party '
    "provides written notice of non-renewal at least ____________ days before the end of the then-current term, "
    "unless earlier terminated as provided in Section 9."
)

add_heading_custom("3. Fees & Payment")
add_body("3.1 Fees. In consideration for the Services, Client shall pay Contractor as follows:")

table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"
headers = ["Service / Milestone", "Fee / Rate", "Due"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for run in p.runs:
            set_run_font(run, bold=True, size=10)
rows_data = [
    ("Initial build / setup (Exhibit A)", "$____________", "Upon signing / milestone schedule"),
    ("Ongoing maintenance & support", "$____________ / month", "Monthly, in advance"),
    ("Hourly rate (out-of-scope work)", "$____________ / hour", "Upon invoice"),
    ("Hosting / third-party fees (if pass-through)", "At cost + ____%", "Upon invoice"),
]
for r, row in enumerate(rows_data, start=1):
    for c, val in enumerate(row):
        cell = table.rows[r].cells[c]
        cell.text = val
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10)

add_blank_line()
add_body(
    "3.2 Invoices & Late Payment. Invoices are due within ____________ days of receipt. Late payments accrue "
    "interest at ____% per month (or the maximum permitted by law, if lower). Contractor may suspend Services for "
    "accounts more than ____________ days past due, upon written notice."
)
add_body(
    "3.3 Third-Party Costs. Client is responsible for all third-party costs incurred on Client's behalf (domain "
    "registration, hosting, cloud databases, licenses, plugins, stock media, payment processor fees, email provider "
    "fees, SMS, analytics, etc.), which shall be itemized and pre-approved by Client except in cases of emergency "
    "remediation reasonably necessary to protect availability, security, or data integrity."
)

add_heading_custom("4. Ownership & Intellectual Property")
add_body(
    "4.1 Work Product. Upon full payment of all applicable fees for the relevant deliverable, Contractor assigns to "
    "Client all right, title, and interest in the custom code, design files, configuration, content, and documentation "
    'created specifically for the Platforms under this Agreement ("Work Product"), excluding Contractor\'s Pre-Existing '
    "IP and any Third-Party Materials."
)
add_body(
    '4.2 Pre-Existing IP. "Pre-Existing IP" means tools, libraries, frameworks, snippets, templates, or methodologies '
    "owned by Contractor prior to, or developed independently of, this engagement. Contractor grants Client a "
    "perpetual, non-exclusive, royalty-free license to use such Pre-Existing IP solely as incorporated into the Platforms."
)
add_body(
    '4.3 Third-Party Materials. "Third-Party Materials" means software, plugins, themes, stock assets, SDKs, or '
    "platforms licensed from third parties (including payment processors, email providers, hosting platforms, and "
    "open-source components). Client's rights to these are governed by the applicable third-party license terms. "
    "Contractor will identify material Third-Party Materials upon request."
)
add_body(
    "4.4 Source Control & Credentials. Work Product includes application source code in version-control repositories "
    "(e.g., GitHub), deployment configurations, and documentation reasonably necessary to operate the Platforms. "
    "Credentials, API keys, and production secrets remain Client Property; Contractor may hold temporary operational "
    "access only as needed to perform Services and shall return or revoke such access on request or termination."
)
add_body(
    "4.5 Portfolio. Contractor retains the right to reference the Platforms in its professional portfolio unless Client "
    "requests otherwise in writing, subject to Section 5 (Confidentiality) and without disclosing patient/client PHI "
    "or Confidential Information."
)

add_heading_custom("5. Confidentiality")
add_body(
    "Each Party may have access to non-public business, technical, financial, patient/client, or proprietary "
    'information of the other Party ("Confidential Information"). Each Party agrees to: (a) use the other\'s '
    "Confidential Information solely to perform its obligations under this Agreement; (b) protect it with at least "
    "the same degree of care it uses for its own confidential information, and no less than reasonable care; and "
    "(c) not disclose it to third parties without prior written consent, except to employees, contractors, or advisors "
    "bound by similar confidentiality obligations, or as required by law. This obligation survives termination of this "
    "Agreement for a period of ____________ years (or indefinitely with respect to trade secrets and Protected Health "
    "Information)."
)

add_heading_custom("6. Data Privacy, Security & HIPAA")
add_body(
    "6.1 Safeguards. Contractor shall implement and maintain reasonable administrative, technical, and physical "
    "safeguards to protect the Platforms and any data collected or processed through them, consistent with industry "
    "standards for similar healthcare-adjacent and e-commerce systems."
)
add_body(
    "6.2 Personal & Health Data. The Platforms collect, store, or process personal information and may collect "
    "health-related or wellness intake data, signatures, payment-related metadata, and clinical-review workflow data. "
    "Contractor shall comply with applicable data protection laws and cooperate with Client's privacy and security "
    "policies."
)
add_body(
    '6.3 HIPAA / BAA. Because Contractor may create, receive, maintain, or transmit Protected Health Information '
    '("PHI") on Client\'s behalf in connection with intakes, charts, portals, or related systems, the Parties shall '
    'execute a separate Business Associate Agreement ("BAA") prior to such access (or promptly after the Effective '
    "Date if such access has already begun). The BAA shall control in the event of conflict with this Agreement "
    "regarding PHI."
)
add_body(
    "6.4 Security Incidents. Contractor shall promptly notify Client of any known or suspected security incident, "
    "unauthorized access, or data breach affecting the Platforms or Client data, and reasonably assist with remediation "
    "and any required notifications."
)
add_body(
    "6.5 Data Ownership. Client owns all Client Data, including patient/client intake submissions, charts, orders, "
    "and operational records stored in Client systems. Contractor may process Client Data solely to perform Services."
)

add_heading_custom("7. Independent Contractor Relationship")
add_body(
    "Contractor is an independent contractor, not an employee, partner, or agent of Client. Contractor is solely "
    "responsible for its own taxes, insurance, benefits, and compliance with applicable employment and tax laws. "
    "Nothing in this Agreement creates a joint venture, partnership, or employment relationship. Contractor has no "
    "authority to bind Client to any obligation."
)

add_heading_custom("8. Warranties")
for item in [
    "Contractor warrants that the Services will be performed in a professional and workmanlike manner consistent with generally accepted industry standards for comparable custom web and application development.",
    "Contractor warrants that, to its knowledge, the Work Product will not infringe the intellectual property rights of any third party.",
    'Contractor will remedy, at no additional charge, any material defects in the Work Product reported by Client within ____________ days of delivery or production deployment of the applicable deliverable ("Warranty Period").',
    'EXCEPT AS EXPRESSLY STATED HEREIN, THE SERVICES AND WORK PRODUCT ARE PROVIDED "AS IS," WITHOUT WARRANTIES OF ANY KIND, WHETHER EXPRESS OR IMPLIED, INCLUDING MERCHANTABILITY OR FITNESS FOR A PARTICULAR PURPOSE. CONTRACTOR DOES NOT WARRANT UNINTERRUPTED OR ERROR-FREE OPERATION OF THIRD-PARTY SERVICES (INCLUDING PAYMENT PROCESSORS, HOSTING PROVIDERS, OR EMAIL PROVIDERS).',
]:
    add_bullet(item)

add_heading_custom("9. Termination")
add_body(
    "9.1 Termination for Convenience. Either Party may terminate this Agreement for any reason upon ____________ "
    "days' written notice."
)
add_body(
    "9.2 Termination for Cause. Either Party may terminate immediately upon written notice if the other Party "
    "materially breaches this Agreement and fails to cure within ____________ days of written notice of such breach "
    "(or immediately if the breach is not reasonably curable)."
)
add_body(
    "9.3 Effect of Termination. Client shall pay for all Services performed and approved expenses incurred through "
    "the effective date of termination. Within ____________ days after termination or Client's written request, "
    "Contractor shall deliver or make available: (a) current Work Product and source repositories; (b) deployment "
    "and environment documentation reasonably needed to operate the Platforms; (c) credentials and access then held "
    "by Contractor (or confirmation of revocation); and (d) a reasonable transition briefing to Client or a successor "
    "contractor. Contractor shall not retain Client production secrets after handover except as required by law or "
    "as mutually agreed in writing for a limited transition period."
)
add_body("9.4 Survival. Sections 4, 5, 6, 10, 11, and 12 survive termination.")

add_heading_custom("10. Limitation of Liability")
add_body(
    "EXCEPT FOR BREACHES OF SECTION 5 (CONFIDENTIALITY), OBLIGATIONS UNDER A BAA WITH RESPECT TO PHI, "
    "INDEMNIFICATION OBLIGATIONS UNDER SECTION 11, OR GROSS NEGLIGENCE / WILLFUL MISCONDUCT, NEITHER PARTY "
    "SHALL BE LIABLE FOR ANY INDIRECT, INCIDENTAL, SPECIAL, OR CONSEQUENTIAL DAMAGES, AND EACH PARTY'S TOTAL "
    "LIABILITY UNDER THIS AGREEMENT SHALL NOT EXCEED THE TOTAL FEES PAID OR PAYABLE BY CLIENT UNDER THIS AGREEMENT "
    "IN THE ____________ MONTHS PRECEDING THE CLAIM."
)

add_heading_custom("11. Indemnification")
add_body(
    "Each Party shall indemnify, defend, and hold harmless the other Party from third-party claims, damages, and "
    "reasonable expenses (including attorneys' fees) arising from: (a) its breach of this Agreement; (b) its gross "
    "negligence or willful misconduct; or (c) its violation of applicable law. This indemnification obligation is "
    "subject to the limitations in Section 10 where applicable, except to the extent a BAA provides otherwise regarding PHI."
)

add_heading_custom("12. General Provisions")
for item in [
    "Governing Law: This Agreement is governed by the laws of the State of ____________, without regard to conflict-of-laws principles.",
    "Dispute Resolution: The Parties shall first attempt to resolve disputes informally in good faith. Unresolved disputes shall be resolved through ____________ (e.g., binding arbitration / courts of ____________ County).",
    "Assignment: Neither Party may assign this Agreement without the other's prior written consent, except in connection with a merger, acquisition, or sale of substantially all assets, provided the assignee assumes this Agreement in writing.",
    "Force Majeure: Neither Party is liable for delays or failures caused by events beyond its reasonable control, including outages of third-party cloud, payment, or communications providers, provided the affected Party gives prompt notice and uses reasonable efforts to mitigate.",
    "Entire Agreement: This Agreement, together with any Exhibits, SOWs, change orders, and any executed BAA, constitutes the entire agreement between the Parties regarding its subject matter and supersedes all prior discussions or agreements regarding that subject matter.",
    "Amendment: This Agreement may only be amended by a written instrument signed by both Parties.",
    "Severability: If any provision is held unenforceable, the remaining provisions remain in full force and effect.",
    "Notices: All notices shall be in writing and delivered by email with confirmation of receipt, overnight courier, or certified mail to the addresses set forth above (or such other address as a Party may designate in writing).",
    "Counterparts: This Agreement may be executed in counterparts (including electronic signature), each of which is deemed an original.",
]:
    add_bullet(item)

add_heading_custom("Signatures")
add_body("IN WITNESS WHEREOF, the Parties have executed this Agreement as of the Effective Date.")
add_blank_line()
add_body("CLIENT — WellnessTech Corporation (and covered affiliates as applicable)", bold=True)
add_body("Signature: ____________________________________")
add_body("Name: ____________________________________")
add_body("Title: ____________________________________")
add_body("Date: ____________________________________")
add_blank_line()
add_body("CONTRACTOR", bold=True)
add_body("Signature: ____________________________________")
add_body("Name: ____________________________________")
add_body("Title: ____________________________________")
add_body("Date: ____________________________________")

doc.add_page_break()
add_title("EXHIBIT A — STATEMENT OF WORK")
add_subtitle("(Initial Scope for WellnessTech Digital Platforms)")

add_body(
    "This Exhibit A is incorporated into the Agreement. Blanks and bracketed items should be completed before signing."
)
add_blank_line()

fields = [
    ("Project Name", "WellnessTech Multi-Brand Digital Platforms — Development & Management"),
    ("Client Entity", "WellnessTech Corporation (parent)"),
    (
        "Sites / Brands Covered",
        "KIAN Privé; KIAN Privé Wellness Hub / Privé Therapeutics; WellnessTech Bio Distribution; Facial Aesthetics / Facial Design Studio; 4EverGlow Wellness; JOA (when developed); other properties listed below: ________________",
    ),
    (
        "Platform / Tech Stack",
        "Custom web applications (e.g., Next.js and/or TanStack Start), APIs, PostgreSQL / Prisma (or successor), hosted cloud deployments, GitHub (or successor) source control — not a single WordPress/Webflow brochure site unless separately specified",
    ),
    (
        "Key Deliverables (built / in scope)",
        "Public brand websites; shop/catalog; booking flows; clinical / wellness intake forms (including peptide / GLP review intakes and related consents); $75 provider review deposit / fee collection where applicable; partner-clinic intake forwarding into central Client systems; practitioner / admin portals; email notifications; EHR / Wellness Tech chart routing as directed by Client; ongoing maintenance as listed below",
    ),
    ("Design Assets / Brand Guidelines Provided by Client", "____________________________________"),
    ("Timeline / Milestones", "____________________________________"),
    ("Repositories / Source Control", "____________________________________ (org/repos; access roles)"),
    ("Hosting / Cloud Provider(s)", "____________________________________"),
    (
        "Third-Party Integrations",
        "Authorize.net (or successor); Resend / transactional email; booking/scheduling; partner intake APIs; analytics; other: ________________",
    ),
    ("Environments", "Production and staging/development as applicable: ________________"),
    (
        "Maintenance Scope (in-scope monthly)",
        "Bug fixes for in-scope features; security/dependency updates; backup verification as configured; uptime troubleshooting; minor content/config changes within agreed monthly hours: ____ hours/month",
    ),
    (
        "Response Time / SLA",
        "Critical (site down, payment failure, intake submission failure): ____ hours response. High: ____ business hours. Normal content/feature requests: queued within maintenance hours or billed hourly",
    ),
    (
        "Out-of-Scope Items",
        "Medical judgment / prescribing; clinical protocol authorship; pharmacy operations; unpaid brand properties not listed above; major new product builds without change order; unpaid third-party license fees; training beyond agreed hours; SEO content marketing campaigns",
    ),
    (
        "Deliverables Acceptance",
        "Client will review milestone deliverables within ____ business days. Silence after that period constitutes acceptance unless Client provides a written defect list",
    ),
    ("Fees Cross-Reference", "As set forth in Section 3 of the Agreement (complete fee table before signing)"),
]

for label, value in fields:
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: ")
    set_run_font(run, bold=True, size=11)
    run2 = p.add_run(value)
    set_run_font(run2, size=11)

add_blank_line()
add_heading_custom("Exhibit A Acknowledgment")
add_body("Client: ___________________________ Date: ____________")
add_body("Contractor: _______________________ Date: ____________")

add_blank_line()
p = doc.add_paragraph()
run = p.add_run(
    "IMPORTANT: This document is a customized business template prepared for discussion of the Parties' "
    "actual multi-site digital platforms. It does not constitute legal advice. Have this Agreement and a "
    "separate HIPAA Business Associate Agreement reviewed by a licensed attorney in your governing jurisdiction before use or signature."
)
set_run_font(run, italic=True, size=10)

out_dir = Path(r"C:\Users\FindMeAnywhere\Desktop\kianprive\docs")
out_dir.mkdir(parents=True, exist_ok=True)
path = out_dir / "WellnessTech_Website_Development_Management_Agreement.docx"
doc.save(path)
print(path)
